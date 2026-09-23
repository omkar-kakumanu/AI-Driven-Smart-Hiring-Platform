import os
import pymupdf as fitz
import docx
import pandas as pd
from resume_parser import (
    process_resume,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_candidate_info,
    generate_profile
)

def create_sample_pdf(pdf_path):
    doc = fitz.open()
    page = doc.new_page()
    sample_text = """
    John Doe
    john.doe@example.com
    +1 555-234-5678

    Education:
    Stanford University

    Experience:
    Tech Corp

    Skills:
    Python, Java, SQL, Machine Learning, TensorFlow, AWS, Docker, Kubernetes

    Certifications:
    AWS Certified Solutions Architect
    """
    page.insert_text((50, 50), sample_text)
    doc.save(pdf_path)
    doc.close()

def create_sample_docx(docx_path):
    doc = docx.Document()
    doc.add_heading("Jane Smith", level=1)
    doc.add_paragraph("Email: jane.smith@techcorp.io")
    doc.add_paragraph("Phone: +1 555-987-6543")
    doc.add_paragraph("Education: MIT Institute of Technology")
    doc.add_paragraph("Experience: NextGen AI Solutions")
    doc.add_paragraph("Skills: Python, SQL, Machine Learning, PyTorch, Docker, AWS, NLP")
    doc.add_paragraph("Certifications: Professional Data Engineer")
    doc.save(docx_path)

def test_pipeline():
    pdf_file = "sample_test_resume.pdf"
    docx_file = "sample_test_resume.docx"

    print("=== Generating Sample PDF Resume ===")
    create_sample_pdf(pdf_file)
    
    print("=== Generating Sample DOCX Resume ===")
    create_sample_docx(docx_file)

    try:
        # Test 1: Direct Candidate Info Extraction from Text
        print("\n--- Testing Candidate Info Extraction (spaCy + Regex + PhraseMatcher) ---")
        sample_text = (
            "Sarah Johnson\n"
            "sarah@example.com\n"
            "(555) 123-4567\n"
            "Stanford University\n"
            "Google\n"
            "Skills: Python, Machine Learning, SQL, PyTorch, AWS\n"
        )
        cand_info = extract_candidate_info(sample_text)
        print("Candidate Info Dict:")
        print(cand_info)
        
        assert cand_info["name"] == "Sarah Johnson"
        assert cand_info["email"] == "sarah@example.com"
        assert "(555) 123-4567" in cand_info["phone"]
        assert "Python" in cand_info["skills"]
        assert "Machine Learning" in cand_info["skills"]
        assert "SQL" in cand_info["skills"]
        assert any("Stanford" in edu for edu in cand_info["education"])
        print("[PASS] Candidate Info extraction passed successfully!")

        # Test 2: PDF Parsing End-to-End
        print("\n--- Testing PDF Resume Extraction (PyMuPDF -> PhraseMatcher -> DataFrame) ---")
        pdf_df = process_resume(pdf_file)
        print("Pandas DataFrame output for PDF:")
        print(pdf_df)
        assert isinstance(pdf_df, pd.DataFrame)
        assert not pdf_df.empty
        pdf_dict = pdf_df.to_dict(orient="records")[0]
        assert pdf_dict["name"] == "John Doe"
        assert pdf_dict["email"] == "john.doe@example.com"
        assert "Python" in pdf_dict["skills"]
        assert "Stanford" in pdf_dict["education"]
        assert "Tech Corp" in pdf_dict["experience"]
        print("[PASS] PDF Extraction test passed successfully!")

        # Test 3: DOCX Parsing End-to-End
        print("\n--- Testing DOCX Resume Extraction (python-docx -> PhraseMatcher -> DataFrame) ---")
        docx_df = process_resume(docx_file)
        print("Pandas DataFrame output for DOCX:")
        print(docx_df)
        assert isinstance(docx_df, pd.DataFrame)
        assert not docx_df.empty
        docx_dict = docx_df.to_dict(orient="records")[0]
        assert docx_dict["name"] == "Jane Smith"
        assert docx_dict["email"] == "jane.smith@techcorp.io"
        assert "PyTorch" in docx_dict["skills"]
        print("[PASS] DOCX Extraction test passed successfully!")

        print("\n=======================================================")
        print("SUCCESS: All Enterprise AI Resume Parser tests verified!")
        print("=======================================================")

    finally:
        if os.path.exists(pdf_file):
            os.remove(pdf_file)
        if os.path.exists(docx_file):
            os.remove(docx_file)

if __name__ == "__main__":
    test_pipeline()
